"""
Estilos, rangos con nombre, bordes y merge.
Ejecutar: python 7_estilos_y_rangos.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.workbook.defined_name import DefinedName


def estilos_y_rangos() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Estilos"

    datos: list[tuple[str, int]] = [("Ana", 85), ("Luis", 92), ("Marta", 78)]

    ws["A1"] = "Nombre"
    ws["B1"] = "Nota"

    for i, (nombre, nota) in enumerate(datos, start=2):
        ws.cell(row=i, column=1, value=nombre)
        ws.cell(row=i, column=2, value=nota)

    # Estilos básicos
    ws["A1"].font = Font(bold=True, italic=True, color="FFFFFF")
    ws["B1"].font = Font(bold=True, color="FFFFFF")

    encabezado_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    ws["A1"].fill = encabezado_fill
    ws["B1"].fill = encabezado_fill

    borde_simple = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    for fila in ws["A1:B4"]:
        for celda in fila:
            celda.border = borde_simple

    # Rango con nombre (A2:B4)
    rango_notas = DefinedName(name="RangoNotas", attr_text="Estilos!$A$2:$B$4")
    wb.defined_names.add(rango_notas)

    # Merge
    ws.merge_cells("D1:E1")
    ws["D1"] = "Resumen"
    ws["D1"].font = Font(bold=True)

    wb.save("estilos_rangos.xlsx")
    print("Archivo estilos_rangos.xlsx guardado.")


if __name__ == "__main__":
    estilos_y_rangos()
