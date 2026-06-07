"""
Leer un archivo Excel y recorrer la zona con datos.
Ejecutar: python 6_leer_archivo.py [ruta]
Si no se indica ruta, usa "ventas.xlsx" por defecto.
"""
import sys
from openpyxl import load_workbook


def leer_archivo(ruta: str = "ventas.xlsx") -> None:
    wb = load_workbook(ruta)

    # Seleccionar hoja por nombre si existe, si no usar la activa
    sheet_name = "Ventas"
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    print("Filas con datos:", ws.max_row)
    print("Columnas con datos:", ws.max_column)

    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column, values_only=True):
        print(fila)


if __name__ == "__main__":
    ruta = sys.argv[1] if len(sys.argv) > 1 else "ventas.xlsx"
    print(f"Leyendo: {ruta}")
    leer_archivo(ruta)
