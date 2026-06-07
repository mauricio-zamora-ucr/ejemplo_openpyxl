"""
Teoría mínima: Workbook, Worksheet, celdas y rangos.
Ejecutar: python 2_teoria_basica.py
"""
from openpyxl import Workbook


def teoria_basica() -> None:
    wb: Workbook = Workbook()
    ws = wb.active
    ws.title = "Datos"

    ws["A1"] = "Nombre"
    ws.cell(row=1, column=2, value="Edad")

    rango = ws["A1:B1"]
    print("Rango objeto:", rango)

    print("Hojas:", wb.sheetnames)
    print("Cantidad de hojas:", len(wb.worksheets))


def hojas() -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Resumen"
    wb.create_sheet("Notas")
    wb.create_sheet("Promedios")

    por_nombre = wb["Notas"]
    por_indice = wb.worksheets[0]

    print("Por nombre:", por_nombre.title)
    print("Por índice:", por_indice.title)


if __name__ == "__main__":
    teoria_basica()
    hojas()
