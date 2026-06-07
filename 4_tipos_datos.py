"""
Tipos de datos básicos: texto, números, booleanos, fechas y moneda.
Ejecutar: python 4_tipos_datos.py
"""
from datetime import datetime
from openpyxl import Workbook


def tipos_datos() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tipos"

    ws["A1"] = "Texto"
    ws["B1"] = "Hola"

    ws["A2"] = "Entero"
    ws["B2"] = 42

    ws["A3"] = "Decimal"
    ws["B3"] = 3.1416

    ws["A4"] = "Booleano"
    ws["B4"] = True

    ws["A5"] = "Fecha"
    ws["B5"] = datetime(2026, 6, 7)
    ws["B5"].number_format = "DD/MM/YYYY"

    ws["A6"] = "Moneda"
    ws["B6"] = 1234.5
    ws["B6"].number_format = '"$"#,##0.00'

    wb.save("tipos.xlsx")
    print("Archivo tipos.xlsx guardado.")


if __name__ == "__main__":
    tipos_datos()
