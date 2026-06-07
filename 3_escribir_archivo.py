"""
Escribir un archivo Excel de ejemplo (`ventas.xlsx`).
Ejecutar: python 3_escribir_archivo.py
"""
from openpyxl import Workbook


def escribir_archivo() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    # Referencia por letra+número
    ws["A1"] = "Producto"
    ws["B1"] = "Cantidad"

    # Referencia por fila/columna
    ws.cell(row=2, column=1, value="Cuaderno")
    ws.cell(row=2, column=2, value=12)

    ws.cell(row=3, column=1, value="Lápiz")
    ws.cell(row=3, column=2, value=20)

    wb.save("ventas.xlsx")
    print("Archivo ventas.xlsx guardado.")


if __name__ == "__main__":
    escribir_archivo()
