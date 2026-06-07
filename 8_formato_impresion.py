"""
Tamaños de columnas/filas, orientación y área de impresión.
Ejecutar: python 8_formato_impresion.py
"""
from openpyxl import Workbook
from openpyxl.worksheet.page import PageMargins


def formato_impresion() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Impresión"

    filas: list[dict[str, str | int]] = [
        {"producto": "Cuaderno universitario", "cantidad": 12},
        {"producto": "Lápiz HB", "cantidad": 20},
        {"producto": "Borrador", "cantidad": 7},
    ]

    ws["A1"] = "Producto"
    ws["B1"] = "Cantidad"

    for i, item in enumerate(filas, start=2):
        ws.cell(row=i, column=1, value=item["producto"])
        ws.cell(row=i, column=2, value=item["cantidad"])

    # Ancho manual de columnas
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12

    # Alto manual de fila
    ws.row_dimensions[1].height = 24

    # Ajuste "automático" básico (calculado)
    for col in ["A", "B"]:
        max_len = 0
        for celda in ws[col]:
            valor = "" if celda.value is None else str(celda.value)
            max_len = max(max_len, len(valor))
        ws.column_dimensions[col].width = max_len + 2

    # Configuración de página
    ws.page_setup.orientation = "landscape"   # horizontal
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.75, bottom=0.75)

    # Área de impresión
    ws.print_area = "A1:B4"

    wb.save("impresion.xlsx")
    print("Archivo impresion.xlsx guardado.")


if __name__ == "__main__":
    formato_impresion()
